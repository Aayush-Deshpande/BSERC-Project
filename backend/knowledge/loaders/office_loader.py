"""
Local Office (DOCX / PPTX) Document Loader for Engine Directives.
DRDO / iDEX Problem Statement ID: 26054

Extracts text from Word documents and PowerPoint presentations on-device.
Zero cloud dependencies.
"""

from pathlib import Path
from typing import Dict, Any, List


def load_docx(file_path: Path) -> List[Dict[str, Any]]:
    """Extracts paragraphs and tables from a Word (.docx) document."""
    path_obj = Path(file_path)
    if not path_obj.exists():
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    try:
        from docx import Document
        doc = Document(str(path_obj))
        text_blocks = []
        for p in doc.paragraphs:
            if p.text.strip():
                text_blocks.append(p.text.strip())
        
        # Also parse table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([c.text.strip() for c in row.cells if c.text.strip()])
                if row_text:
                    text_blocks.append(row_text)

        full_text = "\n\n".join(text_blocks)
        if full_text:
            return [{
                "content": full_text,
                "page_number": 1,
                "source": path_obj.name,
                "type": "docx"
            }]
    except ImportError:
        pass

    return []


def load_pptx(file_path: Path) -> List[Dict[str, Any]]:
    """Extracts slide text from a PowerPoint (.pptx) presentation."""
    path_obj = Path(file_path)
    if not path_obj.exists():
        raise FileNotFoundError(f"PPTX file not found: {file_path}")

    try:
        from pptx import Presentation
        prs = Presentation(str(path_obj))
        slides = []
        for idx, slide in enumerate(prs.slides):
            slide_texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_texts.append(shape.text.strip())
            if slide_texts:
                slides.append({
                    "content": "\n".join(slide_texts),
                    "page_number": idx + 1,
                    "source": path_obj.name,
                    "type": "pptx"
                })
        return slides
    except ImportError:
        pass

    return []
